from __future__ import annotations

import argparse
import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


TITLE_PREFIX = "新能源汽车"
TOP_TOC = [
    ("1 绪论", "4"),
    ("2 关键技术与理论基础", "6"),
    ("3 系统需求分析与总体设计", "7"),
    ("4 关键模块设计与实现", "11"),
    ("5 实验设计与结果分析", "16"),
    ("6 创新点、工作量与不足", "25"),
    ("7 总结与展望", "27"),
    ("参考文献", "28"),
    ("附录 A 图表来源与复现清单", "29"),
    ("附录 B 实验复现实施说明", "30"),
    ("致谢", "31"),
]


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--reference", required=True)
    parser.add_argument("--source", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    reference = Path(args.reference)
    source = Path(args.source)
    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, output)

    ref = Document(reference)
    doc = Document(output)

    apply_section_setup(doc, ref)
    remove_preface_cover_line(doc)
    cleanup_front_matter(doc)
    rebuild_toc(doc)
    restyle_paragraphs(doc)
    restyle_tables(doc)
    apply_footer_page_number(doc)
    doc.save(output)


def apply_section_setup(doc: Document, ref: Document) -> None:
    ref_section = ref.sections[0]
    for section in doc.sections:
        section.page_width = ref_section.page_width
        section.page_height = ref_section.page_height
        section.left_margin = ref_section.left_margin
        section.right_margin = ref_section.right_margin
        section.top_margin = ref_section.top_margin
        section.bottom_margin = ref_section.bottom_margin
        section.header_distance = ref_section.header_distance
        section.footer_distance = ref_section.footer_distance


def remove_preface_cover_line(doc: Document) -> None:
    for p in list(doc.paragraphs):
        if p.text.strip() == "本科毕业论文（设计）正文稿":
            remove_element(p._element)
            break


def rebuild_toc(doc: Document) -> None:
    toc_idx = find_paragraph_index(doc, "目录")
    first_body_idx = next(
        (
            i
            for i, p in enumerate(doc.paragraphs)
            if p.text.strip() == "1 绪论" and is_heading_style_name(p.style.name)
        ),
        None,
    )
    if toc_idx is None or first_body_idx is None:
        return

    body = doc.element.body
    for p in list(doc.paragraphs[toc_idx + 1 : first_body_idx]):
        remove_element(p._element)

    insert_index = list(body).index(doc.paragraphs[toc_idx]._element) + 1
    for text, page in reversed(TOP_TOC):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.first_line_indent = None
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        set_run_font(p.add_run(f"{text}  " + "." * max(8, 40 - len(text)) + f"{page}"), 13)
        body.remove(p._element)
        body.insert(insert_index, p._element)


def cleanup_front_matter(doc: Document) -> None:
    first_body_idx = next(
        (
            i
            for i, p in enumerate(doc.paragraphs)
            if p.text.strip() == "1 绪论" and is_heading_style_name(p.style.name)
        ),
        None,
    )
    if first_body_idx is None:
        return
    for p in list(doc.paragraphs[:first_body_idx]):
        if p.text.strip():
            continue
        remove_element(p._element)


def restyle_paragraphs(doc: Document) -> None:
    first_body_seen = False
    title_done = False
    list_counter = 0
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            list_counter = 0
            continue
        clear_para_formatting(p)
        strip_numbering(p)

        if not title_done and TITLE_PREFIX in text:
            format_title(p)
            title_done = True
            list_counter = 0
            continue

        if text == "摘要":
            format_cn_heading(p)
            list_counter = 0
            continue
        if text == "Abstract":
            format_en_heading(p)
            list_counter = 0
            continue
        if text == "目录":
            format_cn_heading(p)
            list_counter = 0
            continue
        if text == "1 绪论" and not first_body_seen:
            p.paragraph_format.page_break_before = True
            first_body_seen = True
            format_body_heading(p)
            list_counter = 0
            continue
        if is_toc_line(text):
            format_toc_line(p)
            list_counter = 0
            continue
        if text in {"参考文献", "致谢"} or text.startswith("附录"):
            format_body_heading(p)
            p.paragraph_format.page_break_before = text in {"参考文献"} or text.startswith("附录 A")
            list_counter = 0
            continue
        if text.startswith("关键词："):
            format_keyword_line(p, "关键词：")
            list_counter = 0
            continue
        if text.startswith("Key Words:"):
            format_keyword_line(p, "Key Words:")
            list_counter = 0
            continue
        if re.match(r"^\d+\s", text) or re.match(r"^\d+\.\d+", text):
            format_body_heading(p)
            list_counter = 0
            continue
        if text.startswith("图") or text.startswith("表"):
            format_caption(p)
            list_counter = 0
            continue
        if p.style.name.startswith("List"):
            list_counter += 1
            convert_list_to_manual_text(p, list_counter)
            format_body(p)
            continue
        list_counter = 0
        format_body(p)

    apply_page_breaks(doc)


def restyle_tables(doc: Document) -> None:
    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        for ri, row in enumerate(table.rows):
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                set_cell_margins(cell, top=90, start=110, bottom=90, end=110)
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if ri else WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.first_line_indent = None
                    p.paragraph_format.left_indent = None
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
                    for run in p.runs:
                        set_run_font(run, 12, bold=ri == 0)
                        clear_run_color(run)
                shade_cell(cell, "D9E2F3" if ri == 0 else "FFFFFF")


def apply_footer_page_number(doc: Document) -> None:
    for section in doc.sections:
        footer = section.footer
        p = footer.paragraphs[0]
        clear_paragraph(p)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), "PAGE")
        run = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        rfonts = OxmlElement("w:rFonts")
        rfonts.set(qn("w:ascii"), "Times New Roman")
        rfonts.set(qn("w:hAnsi"), "Times New Roman")
        rfonts.set(qn("w:eastAsia"), "宋体")
        rpr.append(rfonts)
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), "24")
        rpr.append(sz)
        run.append(rpr)
        t = OxmlElement("w:t")
        t.text = "1"
        run.append(t)
        fld.append(run)
        p._p.append(fld)


def format_title(p) -> None:
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = None
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.space_before = Pt(24)
    p.paragraph_format.space_after = Pt(28)
    for run in p.runs:
        set_run_font(run, 15, bold=True)
        clear_run_color(run)


def format_cn_heading(p) -> None:
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = None
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    for run in p.runs:
        set_run_font(run, 15, bold=True)
        clear_run_color(run)


def format_en_heading(p) -> None:
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = None
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    for run in p.runs:
        set_run_font(run, 15, bold=False)
        clear_run_color(run)


def format_body_heading(p) -> None:
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = None
    p.paragraph_format.left_indent = None
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.5
    for run in p.runs:
        set_run_font(run, 15, bold=False)
        clear_run_color(run)


def format_body(p) -> None:
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if needs_left_align(p.text.strip()) else WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Pt(28)
    p.paragraph_format.left_indent = None
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.5
    for run in p.runs:
        set_run_font(run, 14, bold=False)
        clear_run_color(run)


def format_caption(p) -> None:
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = None
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    for run in p.runs:
        set_run_font(run, 14, bold=False)
        clear_run_color(run)


def format_toc_line(p) -> None:
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = None
    p.paragraph_format.left_indent = None
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.2
    for run in p.runs:
        set_run_font(run, 13, bold=False)
        clear_run_color(run)


def format_keyword_line(p, prefix: str) -> None:
    text = p.text.strip()
    remainder = text[len(prefix) :].lstrip()
    clear_paragraph(p)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = None
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.5
    set_run_font(p.add_run(prefix), 14, bold=True)
    set_run_font(p.add_run(" " + remainder), 14, bold=False)


def convert_list_to_manual_text(p, number: int) -> None:
    text = p.text.strip()
    if not text:
        return
    if re.match(r"^[（(]\d+[)）]", text) or text.startswith("•") or text.startswith("·"):
        return
    prefix = f"（{number}）"
    clear_paragraph(p)
    p.add_run(prefix + text)


def apply_page_breaks(doc: Document) -> None:
    seen_abstract = seen_abstract_en = seen_toc = False
    for p in doc.paragraphs:
        text = p.text.strip()
        if text == "摘要":
            seen_abstract = True
            continue
        if text == "Abstract":
            p.paragraph_format.page_break_before = True
            seen_abstract_en = True
            continue
        if text == "目录":
            p.paragraph_format.page_break_before = True
            seen_toc = True
            continue
        if text == "1 绪论" and seen_toc:
            p.paragraph_format.page_break_before = True
            continue
        if text in {"参考文献", "附录 A 图表来源与复现清单", "致谢"}:
            p.paragraph_format.page_break_before = True


def clear_para_formatting(p) -> None:
    pf = p.paragraph_format
    pf.first_line_indent = None
    pf.left_indent = None
    pf.right_indent = None
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.5
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.page_break_before = False
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def clear_run_color(run) -> None:
    run.font.color.rgb = RGBColor(0, 0, 0)


def strip_numbering(p) -> None:
    pPr = p._element.pPr
    if pPr is None:
        return
    numPr = pPr.find(qn("w:numPr"))
    if numPr is not None:
        pPr.remove(numPr)


def clear_paragraph(p) -> None:
    for child in list(p._p):
        p._p.remove(child)


def set_run_font(run, size_pt: float, bold: bool = False) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = False
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")
    rfonts.set(qn("w:eastAsia"), "宋体")


def set_cell_margins(cell, top=80, start=80, bottom=80, end=80):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for name, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tcMar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tcMar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)


def is_toc_line(text: str) -> bool:
    return "." in text and any(text.startswith(item[0]) for item in TOP_TOC)


def needs_left_align(text: str) -> bool:
    ascii_chars = sum(1 for ch in text if ch.isascii() and ch.isalnum())
    signals = ["/", "_", "API", "OCR", "BGE", "vLLM", "Next.js", "FastAPI", "pgvector", "Docker", "TypeScript"]
    return ascii_chars >= 18 or any(token in text for token in signals)


def find_paragraph_index(doc: Document, text: str) -> int | None:
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() == text:
            return i
    return None


def is_heading_style_name(name: str) -> bool:
    return "Heading" in name or "标题" in name


def remove_element(element) -> None:
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


if __name__ == "__main__":
    main()
