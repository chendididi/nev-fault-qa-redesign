from __future__ import annotations

import argparse
import shutil
import textwrap
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


FIGURE_SPECS = {
    "图3-1": "fig3_1_architecture.png",
    "图3-2": "fig3_2_model_routing.png",
    "图3-3": "fig3_3_deployment.png",
    "图4-1": "fig4_1_rag_flow.png",
    "图4-2": "fig4_2_ui_overview.png",
    "图5-7": "fig5_7_case_flow.png",
}


TITLE = "新能源汽车多模态诊断问答系统设计与实现"
CN_ABSTRACT = (
    "针对新能源汽车维修资料分散、跨系统诊断复杂和高压安全要求高等问题，本文设计并实现一套多模态可追溯检索增强问答系统。"
    "系统采用 Next.js、FastAPI、PostgreSQL/pgvector、Redis 和 MinIO 构建，支持资料入库、OCR、领域元数据抽取、混合检索、"
    "BGE 重排序与结构化诊断生成。基于 64 条模拟维修样例的实验表明，Hybrid+Rerank 方案在 Hit@1 和 MRR 上均为 1.0000，"
    "引用准确率为 0.9792，安全证据覆盖率为 1.0000。结果说明该系统能在保证证据可追溯的前提下提升维修资料定位和诊断辅助质量。"
)
EN_ABSTRACT = (
    "To address fragmented maintenance documents, cross-system diagnostic complexity and high-voltage safety requirements in new energy vehicle repair, "
    "this thesis designs and implements a multimodal traceable retrieval-augmented question answering system. The system is built with Next.js, FastAPI, "
    "PostgreSQL/pgvector, Redis and MinIO. It supports document ingestion, OCR, domain metadata extraction, hybrid recall, BGE reranking, structured "
    "diagnostic generation and citation-chain presentation. A task-level model routing mechanism connects vLLM, local models and remote OpenAI-compatible "
    "services. Experiments on 64 simulated maintenance diagnosis cases show that the Hybrid+Rerank setting achieves 1.0000 Hit@1 and MRR, 0.9792 citation "
    "accuracy and 1.0000 safety evidence coverage. The results indicate that the system improves maintenance evidence retrieval and diagnostic assistance "
    "while keeping answers traceable."
)


REPLACEMENTS = {
    "面向新能源汽车故障诊断的部署感知多模态检索增强问答系统设计与实现": TITLE,
    "面向新能源汽车故障诊断的 部署感知多模态检索增强问答系统设计与实现": TITLE,
    "关键词：新能源汽车故障诊断；多模态 RAG；混合检索；引用溯源；模型路由；pgvector；vLLM": "关键词：新能源汽车故障诊断；多模态 RAG；混合检索；引用溯源；模型路由",
    "Key Words: New Energy Vehicle; Fault Diagnosis; Multimodal RAG; Hybrid Retrieval; Citation Traceability; Model Routing; pgvector; vLLM": "Key Words: New Energy Vehicle; Multimodal RAG; Hybrid Retrieval; Citation Traceability; Model Routing",
    "5.6 仍需补做的人工评分实验": "5.6 人工评分实验设计",
    "图4-2 系统界面截图组": "图4-2 系统主要功能界面示意图",
    "图5-7 快充互锁异常端到端诊断案例截图": "图5-7 快充互锁异常端到端诊断流程示意图",
    "附录 A 图表补充清单": "附录 A 图表来源与复现清单",
}


PARAGRAPH_REPLACEMENTS = {
    "本科毕业设计（论文）修订稿": "本科毕业论文（设计）正文稿",
    "现有智能问答研究大致经历了关键词检索、语义向量检索、知识图谱问答和检索增强生成等阶段。关键词检索在故障码、标准值、部件编号等精确匹配任务中稳定，但难以处理“快充枪已锁止但 SOC 不上升”这类自然语言症状。向量检索能捕获语义相似性，却可能遗漏故障码、单位、阈值等精确术语。重排序模型通过 query-passage 联合编码提升证据排序质量，但会带来额外延迟。RAG 通过“先检索、后生成”降低大模型幻觉，是当前知识密集型问答的重要范式。":
        "现有智能问答研究大致经历了关键词检索、语义向量检索、知识图谱问答和检索增强生成等阶段。关键词检索在故障码、标准值、部件编号等精确匹配任务中稳定，BM25 等概率检索模型仍是重要基线[2]；密集向量检索能捕获语义相似性，但可能遗漏故障码、单位、阈值等精确术语[7]。重排序模型通过 query-passage 联合编码提升证据排序质量[4]。RAG 通过“先检索、后生成”降低大模型幻觉，是当前知识密集型问答的重要范式[1]。",
    "检索增强生成（Retrieval-Augmented Generation, RAG）的基本思想是在生成回答前从外部知识库检索相关证据，再将证据片段组织进提示词，使大模型基于可追溯资料生成答案。与直接生成相比，RAG 能显著降低无依据编造风险，并能在回答中呈现引用来源。对新能源汽车维修而言，RAG 的关键价值不只是“让模型知道更多知识”，而是把维修结论绑定到具体资料页码、章节、故障码和系统分类，从而满足现场排查的可解释性和安全性要求。":
        "检索增强生成（Retrieval-Augmented Generation, RAG）的基本思想是在生成回答前从外部知识库检索相关证据，再将证据片段组织进提示词，使大模型基于可追溯资料生成答案[1]。与直接生成相比，RAG 能显著降低无依据编造风险，并能在回答中呈现引用来源。对新能源汽车维修而言，RAG 的关键价值不只是“让模型知道更多知识”，而是把维修结论绑定到具体资料页码、章节、故障码和系统分类，从而满足现场排查的可解释性和安全性要求。",
    "当前项目采用 PostgreSQL + pgvector 统一保存业务表、知识分块和向量索引。与额外部署独立向量数据库相比，该方案降低了部署复杂度，并便于将文档、chunk、会话、检索命中和引用元数据保存在同一事务体系中。系统在 chunks 表中保存 content、metadata 和 embedding 字段，使用 pgvector 的余弦距离进行 Top-K 语义召回。metadata 中记录标题、文件名、chunk_index":
        "当前项目采用 PostgreSQL + pgvector 统一保存业务表、知识分块和向量索引[8-9]。与额外部署独立向量数据库相比，该方案降低了部署复杂度，并便于将文档、chunk、会话、检索命中和引用元数据保存在同一事务体系中。系统在 chunks 表中保存 content、metadata 和 embedding 字段，使用 pgvector 的余弦距离进行 Top-K 语义召回。metadata 中记录标题、文件名、chunk_index、页码、章节、系统和故障码等信息。",
    "单纯向量召回对语义相似问题有效，但对于 P1A0C、P1A11、100kΩ、CAN-H、HVIL、READY 等领域词可能出现召回不稳定。系统通过正则表达式和新能源汽车领域词表抽取 DTC、系统名、安全词和诊断术语，再在 chunk 内容和 metadata 中进行关键词补召回。向量召回与关键词召回合并后，先按向量分和关键词命中数量形成候选，再通过本地 BGE reranker 对 query-passage 对进行精排。该设计兼顾了语":
        "单纯向量召回对语义相似问题有效，但对于 P1A0C、P1A11、100kΩ、CAN-H、HVIL、READY 等领域词可能出现召回不稳定。系统通过正则表达式和新能源汽车领域词表抽取 DTC、系统名、安全词和诊断术语，再在 chunk 内容和 metadata 中进行关键词补召回。向量召回与关键词召回合并后，先按向量分和关键词命中数量形成候选，再通过本地 BGE reranker 对 query-passage 对进行精排[5]。该设计兼顾了语义泛化能力和维修术语精确命中能力。",
    "当前技术栈的重要改革是将模型配置从单一 base_url 模式扩展为“模型端点 + 任务路由”拓扑。端点描述 vLLM、Ollama、llama.cpp、LM Studio、第三方 OpenAI-compatible API、本地模型或 Demo Cache；路由描述 chat、vision_ocr、embedding、rerank、fallback_chat 分别使用哪个端点和模型。该设计使聊天与图片理解可以使用学校 4090 服务器":
        "当前技术栈的重要改革是将模型配置从单一 base_url 模式扩展为“模型端点 + 任务路由”拓扑。端点描述 vLLM、Ollama、llama.cpp、LM Studio、第三方 OpenAI-compatible API、本地模型或 Demo Cache；路由描述 chat、vision_ocr、embedding、rerank、fallback_chat 分别使用哪个端点和模型。该设计使聊天与图片理解可以使用学校 4090 服务器上的 vLLM/Qwen2.5-VL 服务[10-11]，embedding 与 rerank 保持本地 BGE 模型，从而兼顾效果、成本和部署稳定性。",
    "现有实验已经覆盖检索侧指标，但为了让论文更像完整学术研究，建议补做 30 条问答结果的人工评分。评分不应凭主观印象，而应按答案正确、引用匹配、安全合规、无幻觉、步骤可执行五个维度进行 0/1 标注。该实验可以证明系统不仅能找对证据，也能把证据组织成符合维修逻辑的回答。":
        "除自动检索指标外，本文设计 30 条问答结果的人工评分方案，用于评价答案是否能把证据组织成符合维修逻辑的诊断建议。评分不凭主观印象，而按答案正确、引用匹配、安全合规、无幻觉、步骤可执行五个维度进行 0/1 标注。该方案可作为后续答辩验收和系统迭代的补充证据。",
    "本附录列出提交前必须补齐或建议补齐的图片。实验类图片已经从当前项目 docs/paper/figures 插入正文；缺少的是系统架构、部署拓扑和真实界面截图，这些需要根据当前运行系统补拍或绘制。":
        "本附录列出正文图表的来源与复现方式。系统架构、模型路由、部署拓扑、RAG 流程和案例流程图依据当前项目代码与部署文件绘制；实验图来自 docs/paper/figures 和 apps/api 的评测输出。若提交最终版，可在此基础上补充真实运行界面截图。",
}


REFERENCES = [
    "[1] Lewis P, Perez E, Piktus A, et al. Retrieval-Augmented Generation for Knowledge-Intensive NLP Tasks[C]//Advances in Neural Information Processing Systems. 2020.",
    "[2] Robertson S, Zaragoza H. The Probabilistic Relevance Framework: BM25 and Beyond[J]. Foundations and Trends in Information Retrieval, 2009, 3(4): 333-389.",
    "[3] Cormack G V, Clarke C L A, Buettcher S. Reciprocal Rank Fusion Outperforms Condorcet and Individual Rank Learning Methods[C]//SIGIR. 2009: 758-759.",
    "[4] Nogueira R, Cho K. Passage Re-ranking with BERT[EB/OL]. arXiv:1901.04085, 2019.",
    "[5] Xiao S, Liu Z, Zhang P, et al. C-Pack: Packaged Resources To Advance General Chinese Embedding[EB/OL]. arXiv preprint, 2023.",
    "[6] Johnson J, Douze M, Jegou H. Billion-scale Similarity Search with GPUs[J]. IEEE Transactions on Big Data, 2019, 7(3): 535-547.",
    "[7] Karpukhin V, Oguz B, Min S, et al. Dense Passage Retrieval for Open-Domain Question Answering[C]//EMNLP. 2020: 6769-6781.",
    "[8] PostgreSQL Global Development Group. PostgreSQL Documentation[EB/OL]. https://www.postgresql.org/docs/.",
    "[9] pgvector Contributors. pgvector: Open-source Vector Similarity Search for Postgres[EB/OL]. https://github.com/pgvector/pgvector.",
    "[10] Kwon W, Li Z, Zhuang S, et al. Efficient Memory Management for Large Language Model Serving with PagedAttention[C]//SOSP. 2023.",
    "[11] Qwen Team. Qwen2.5-VL Technical Report[EB/OL]. 2025.",
    "[12] 国家市场监督管理总局, 国家标准化管理委员会. GB 18384-2020 电动汽车安全要求[S]. 北京: 中国标准出版社, 2020.",
]


TABLE_CAPTIONS = {
    "角色": "表3-1 角色功能与论文关注点",
    "数据表": "表3-2 核心数据表设计",
    "步骤": "表4-1 混合检索与重排序步骤",
    "生成模板约束": "表4-2 结构化生成模板约束",
    "实现文件": "表4-3 系统主要实现文件",
    "项目": "表5-1 实验环境与配置",
    "指标": "表5-2 自动评测指标说明",
    "方法": "表5-3 检索消融实验结果",
    "环节": "表5-4 快充互锁异常案例分析",
    "系统": "表5-5 人工评分维度与记录方式",
    "工作类别": "表6-1 项目实现工作量说明",
    "编号": "表A-1 图表来源与复现说明",
    "命令": "表B-1 检索评测复现命令",
}

TOC_LINES = [
    "1 绪论........................................4",
    "2 关键技术与理论基础..........................6",
    "3 系统需求分析与总体设计......................7",
    "4 关键模块设计与实现.........................10",
    "5 实验设计与结果分析.........................14",
    "6 创新点、工作量与不足.......................21",
    "7 总结与展望.................................22",
    "参考文献....................................23",
    "附录 A 图表来源与复现清单....................24",
    "附录 B 实验复现实施说明......................24",
    "致谢........................................25",
]


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--workdir", required=True)
    args = parser.parse_args()

    input_path = Path(args.input)
    output_path = Path(args.output)
    workdir = Path(args.workdir)
    figure_dir = workdir / "generated_figures"
    figure_dir.mkdir(parents=True, exist_ok=True)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    figure_paths = generate_figures(figure_dir)
    shutil.copy2(input_path, output_path)
    doc = Document(output_path)

    normalize_document(doc)
    remove_internal_tables(doc)
    replace_paragraphs(doc)
    rebuild_toc(doc)
    replace_placeholder_figures(doc, figure_paths)
    rewrite_supporting_tables(doc)
    add_table_captions(doc)
    rewrite_references(doc)
    add_acknowledgement(doc)
    configure_footer_page_numbers(doc)
    enforce_page_breaks(doc)
    doc.save(output_path)


def generate_figures(out_dir: Path) -> dict[str, Path]:
    font_regular = font_path("msyh.ttc")
    font_bold = font_path("msyhbd.ttc") or font_regular
    paths = {}
    paths["图3-1"] = draw_architecture(out_dir / FIGURE_SPECS["图3-1"], font_regular, font_bold)
    paths["图3-2"] = draw_model_routing(out_dir / FIGURE_SPECS["图3-2"], font_regular, font_bold)
    paths["图3-3"] = draw_deployment(out_dir / FIGURE_SPECS["图3-3"], font_regular, font_bold)
    paths["图4-1"] = draw_rag_flow(out_dir / FIGURE_SPECS["图4-1"], font_regular, font_bold)
    paths["图4-2"] = draw_ui_overview(out_dir / FIGURE_SPECS["图4-2"], font_regular, font_bold)
    paths["图5-7"] = draw_case_flow(out_dir / FIGURE_SPECS["图5-7"], font_regular, font_bold)
    return paths


def font_path(name: str) -> str:
    candidates = [
        Path("C:/Windows/Fonts") / name,
        Path("/mnt/c/Windows/Fonts") / name,
        Path("/usr/share/fonts/truetype/wqy/wqy-microhei.ttc"),
    ]
    for path in candidates:
        if path.exists():
            return str(path)
    return ""


def load_font(path: str, size: int) -> ImageFont.FreeTypeFont:
    if path:
        return ImageFont.truetype(path, size=size)
    return ImageFont.load_default()


def new_canvas(w: int = 2200, h: int = 1300):
    img = Image.new("RGB", (w, h), "white")
    draw = ImageDraw.Draw(img)
    return img, draw


def draw_text(draw, xy, text, font, fill="#111827", anchor=None, max_width=None, line_spacing=8, align="center"):
    if max_width:
        text = wrap_by_pixel(draw, text, font, max_width)
    draw.multiline_text(xy, text, font=font, fill=fill, anchor=anchor, spacing=line_spacing, align=align)


def wrap_by_pixel(draw, text: str, font, max_width: int) -> str:
    lines = []
    for raw_line in text.split("\n"):
        current = ""
        for ch in raw_line:
            test = current + ch
            bbox = draw.textbbox((0, 0), test, font=font)
            if bbox[2] - bbox[0] > max_width and current:
                lines.append(current)
                current = ch
            else:
                current = test
        if current:
            lines.append(current)
    return "\n".join(lines)


def box(draw, xy, title, body, title_font, body_font, *, fill="#F8FAFC", outline="#CBD5E1", accent="#2563EB"):
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=18, fill=fill, outline=outline, width=3)
    draw.rectangle((x1, y1, x1 + 14, y2), fill=accent)
    draw_text(draw, ((x1 + x2) // 2, y1 + 42), title, title_font, anchor="mm", max_width=x2 - x1 - 70)
    if body:
        draw_text(draw, ((x1 + x2) // 2, y1 + 115), body, body_font, fill="#475569", anchor="ma", max_width=x2 - x1 - 70, line_spacing=10)


def arrow(draw, start, end, color="#64748B", width=5):
    draw.line((start, end), fill=color, width=width)
    sx, sy = start
    ex, ey = end
    if abs(ex - sx) >= abs(ey - sy):
        direction = 1 if ex >= sx else -1
        pts = [(ex, ey), (ex - 20 * direction, ey - 12), (ex - 20 * direction, ey + 12)]
    else:
        direction = 1 if ey >= sy else -1
        pts = [(ex, ey), (ex - 12, ey - 20 * direction), (ex + 12, ey - 20 * direction)]
    draw.polygon(pts, fill=color)


def draw_architecture(path: Path, regular: str, bold: str) -> Path:
    img, draw = new_canvas()
    title_font = load_font(bold, 44)
    h_font = load_font(bold, 34)
    b_font = load_font(regular, 26)
    draw_text(draw, (1100, 70), "系统总体架构", title_font, anchor="mm")
    box(draw, (110, 170, 540, 360), "用户层", "管理员：资料入库、模型配置\n维修技师：文本/图片问答、证据查看", h_font, b_font, accent="#DC2626")
    box(draw, (700, 170, 1130, 360), "前端层", "Next.js + TypeScript\n诊断工作台、知识库、模型设置", h_font, b_font, accent="#0F766E")
    box(draw, (1290, 170, 1720, 360), "接口层", "FastAPI\n认证、SSE、RAG、模型测试", h_font, b_font, accent="#2563EB")
    box(draw, (330, 520, 760, 720), "知识工程", "文档上传、OCR、切片\nmetadata 抽取、向量写入", h_font, b_font, accent="#7C3AED")
    box(draw, (890, 520, 1310, 720), "检索生成", "pgvector 召回 + 关键词补召回\nBGE rerank + 结构化生成", h_font, b_font, accent="#EA580C")
    box(draw, (1440, 520, 1870, 720), "模型路由", "chat / vision_ocr / embedding\nrerank / fallback_chat", h_font, b_font, accent="#0891B2")
    box(draw, (170, 890, 560, 1100), "PostgreSQL + pgvector", "业务表、chunks、embedding\nretrieval_hits 引用证据链", h_font, b_font, accent="#2563EB")
    box(draw, (700, 890, 1030, 1100), "Redis", "任务队列\n缓存状态", h_font, b_font, accent="#DC2626")
    box(draw, (1170, 890, 1500, 1100), "MinIO", "原始文档\n图片对象", h_font, b_font, accent="#16A34A")
    box(draw, (1640, 890, 2030, 1100), "模型服务", "本地 BGE、vLLM/Qwen\n远端 API、Demo Cache", h_font, b_font, accent="#9333EA")
    for s, e in [((540, 265), (700, 265)), ((1130, 265), (1290, 265)), ((1505, 360), (1505, 520)), ((540, 720), (365, 890)), ((1100, 720), (365, 890)), ((1100, 720), (865, 890)), ((540, 720), (1335, 890)), ((1655, 720), (1835, 890))]:
        arrow(draw, s, e)
    img.save(path)
    return path


def draw_model_routing(path: Path, regular: str, bold: str) -> Path:
    img, draw = new_canvas()
    title_font = load_font(bold, 44)
    h_font = load_font(bold, 32)
    b_font = load_font(regular, 25)
    draw_text(draw, (1100, 70), "模型端点与任务路由拓扑", title_font, anchor="mm")
    tasks = [
        ("chat", "结构化诊断回答"),
        ("vision_ocr", "仪表/诊断仪图片理解"),
        ("embedding", "query 与 chunk 向量化"),
        ("rerank", "query-passage 精排"),
        ("fallback_chat", "真实模型失败后兜底"),
    ]
    y = 170
    for name, desc in tasks:
        box(draw, (130, y, 520, y + 140), name, desc, h_font, b_font, accent="#2563EB")
        arrow(draw, (520, y + 70), (790, y + 70))
        y += 185
    box(draw, (790, 260, 1160, 800), "model_routes", "task\nendpoint_id\nmodel_name\ntemperature\nmax_tokens\nenabled", h_font, b_font, accent="#EA580C")
    endpoints = [
        ("vLLM / Qwen2.5-VL", "chat、vision_ocr\n学校 4090 GPU"),
        ("Local BGE", "embedding、rerank\nAPI 容器内本地模型"),
        ("Remote API", "OpenAI-compatible\n远端模型备选"),
        ("Demo Cache", "演示兜底\n显式标注演示结果"),
    ]
    y = 170
    for name, desc in endpoints:
        box(draw, (1430, y, 2020, y + 155), name, desc, h_font, b_font, accent="#16A34A")
        arrow(draw, (1160, 530), (1430, y + 78))
        y += 220
    draw_text(draw, (1100, 1135), "核心思想：任务与端点解耦，聊天、视觉、向量化、重排序和兜底可独立迁移。", b_font, fill="#334155", anchor="mm")
    img.save(path)
    return path


def draw_deployment(path: Path, regular: str, bold: str) -> Path:
    img, draw = new_canvas()
    title_font = load_font(bold, 44)
    h_font = load_font(bold, 32)
    b_font = load_font(regular, 25)
    draw_text(draw, (1100, 70), "学校服务器 Docker Compose 部署拓扑", title_font, anchor="mm")
    box(draw, (120, 220, 460, 390), "浏览器", "localhost:13001\n或服务器 Web 端口", h_font, b_font, accent="#DC2626")
    box(draw, (650, 180, 1010, 430), "web 容器", "Next.js 前端\nWEB_PORT=3001\nAPI Base -> api:8000", h_font, b_font, accent="#0F766E")
    box(draw, (1190, 180, 1560, 430), "api 容器", "FastAPI\n/api/chat/stream\n模型测试、RAG 服务", h_font, b_font, accent="#2563EB")
    box(draw, (1750, 180, 2090, 430), "vLLM 容器", "Qwen2.5-VL-7B\nOpenAI-compatible\nVLLM_PORT=8008", h_font, b_font, accent="#7C3AED")
    box(draw, (310, 710, 650, 910), "worker", "异步入库\n解析、OCR、embedding", h_font, b_font, accent="#EA580C")
    box(draw, (820, 700, 1170, 920), "PostgreSQL", "pgvector\nusers / chunks\nretrieval_hits", h_font, b_font, accent="#2563EB")
    box(draw, (1320, 700, 1640, 920), "Redis", "任务队列\n状态缓存", h_font, b_font, accent="#DC2626")
    box(draw, (1780, 700, 2090, 920), "MinIO", "上传文件\n图片对象", h_font, b_font, accent="#16A34A")
    for s, e in [((460, 305), (650, 305)), ((1010, 305), (1190, 305)), ((1560, 305), (1750, 305)), ((1375, 430), (990, 700)), ((1375, 430), (1480, 700)), ((1375, 430), (1935, 700)), ((1375, 430), (480, 710))]:
        arrow(draw, s, e)
    draw_text(draw, (1100, 1125), "deploy-school.sh 固定 Compose 项目名 nev-fault-qa-redesign，配合 export/import 脚本迁移演示数据。", b_font, fill="#334155", anchor="mm")
    img.save(path)
    return path


def draw_rag_flow(path: Path, regular: str, bold: str) -> Path:
    img, draw = new_canvas(2400, 1300)
    title_font = load_font(bold, 44)
    h_font = load_font(bold, 29)
    b_font = load_font(regular, 23)
    draw_text(draw, (1200, 70), "多模态 RAG 问答流程", title_font, anchor="mm")
    row1 = [
        ("文本问题", "DTC、症状、追问"),
        ("图片输入", "仪表/诊断仪照片"),
        ("历史上下文", "最近多轮用户问题"),
        ("query 构造", "文本 + OCR + 上下文"),
    ]
    xs = [90, 650, 1210, 1770]
    for x, (t, b) in zip(xs, row1):
        box(draw, (x, 190, x + 420, 360), t, b, h_font, b_font, accent="#2563EB")
    arrow(draw, (510, 275), (650, 275))
    arrow(draw, (1070, 275), (1210, 275))
    arrow(draw, (1630, 275), (1770, 275))
    row2 = [
        ("向量召回", "pgvector Top-K"),
        ("关键词补召回", "DTC / 系统 / 安全词"),
        ("候选合并", "vector_score + keyword_score"),
        ("BGE rerank", "query-passage 精排"),
    ]
    for x, (t, b) in zip(xs, row2):
        box(draw, (x, 570, x + 420, 740), t, b, h_font, b_font, accent="#EA580C")
    arrow(draw, (1980, 360), (300, 570))
    for i in range(3):
        arrow(draw, (xs[i] + 420, 655), (xs[i + 1], 655))
    row3 = [
        ("Prompt 构造", "编号证据 + 安全约束"),
        ("模型生成", "chat 路由 / fallback"),
        ("结构化回答", "原因、步骤、安全、引用"),
        ("证据持久化", "retrieval_hits 可复核"),
    ]
    for x, (t, b) in zip(xs, row3):
        box(draw, (x, 940, x + 420, 1110), t, b, h_font, b_font, accent="#16A34A")
    arrow(draw, (1980, 740), (300, 940))
    for i in range(3):
        arrow(draw, (xs[i] + 420, 1025), (xs[i + 1], 1025))
    img.save(path)
    return path


def draw_ui_overview(path: Path, regular: str, bold: str) -> Path:
    img, draw = new_canvas()
    title_font = load_font(bold, 44)
    h_font = load_font(bold, 32)
    b_font = load_font(regular, 24)
    draw_text(draw, (1100, 70), "系统主要功能界面示意", title_font, anchor="mm")
    panels = [
        ((120, 180, 980, 520), "知识库管理", "上传 PDF / Word / Excel / 图片\n展示 pending、processing、ready、failed\n错误信息用于定位解析问题", "#0F766E"),
        ((1220, 180, 2080, 520), "诊断工作台", "选择知识库并输入故障现象\n支持图片上传、多轮追问和 SSE 流式回答\n回答按固定诊断结构展开", "#2563EB"),
        ((120, 720, 980, 1060), "证据链展开", "展示来源文件、页码、章节、系统、DTC\n显示 vector_score、rerank_score、关键词命中\n用于复核答案依据", "#EA580C"),
        ((1220, 720, 2080, 1060), "模型配置", "端点管理：vLLM、远端 API、本地模型、Demo Cache\n任务路由：chat、vision_ocr、embedding、rerank\n模型测试返回连通性和耗时", "#7C3AED"),
    ]
    for xy, title, body, accent in panels:
        box(draw, xy, title, body, h_font, b_font, accent=accent)
    arrow(draw, (980, 350), (1220, 350))
    arrow(draw, (1530, 520), (650, 720))
    arrow(draw, (980, 890), (1220, 890))
    draw_text(draw, (1100, 1160), "该图概括当前前端的关键交互，不替代最终提交前可补拍的真实运行截图。", b_font, fill="#334155", anchor="mm")
    img.save(path)
    return path


def draw_case_flow(path: Path, regular: str, bold: str) -> Path:
    img, draw = new_canvas(2300, 1200)
    title_font = load_font(bold, 42)
    h_font = load_font(bold, 29)
    b_font = load_font(regular, 22)
    draw_text(draw, (1150, 65), "P1A0C 快充互锁异常端到端诊断流程", title_font, anchor="mm")
    steps = [
        ("用户问题", "P1A0C 快充互锁异常怎么排查？"),
        ("检索词抽取", "P1A0C、快充、互锁、HVIL"),
        ("混合召回", "向量召回快充/慢充资料\n关键词补召回 DTC 与互锁证据"),
        ("重排序首位证据", "快充慢充系统维修诊断手册\nvector_score=0.7943\nrerank_score=0.9982"),
        ("结构化回答", "可能原因、检查步骤\n高压安全提醒、维修建议\n引用来源与置信提示"),
    ]
    xs = [70, 500, 930, 1360, 1790]
    for x, (t, b) in zip(xs, steps):
        box(draw, (x, 210, x + 370, 570), t, b, h_font, b_font, accent="#2563EB")
    for i in range(4):
        arrow(draw, (xs[i] + 370, 390), (xs[i + 1], 390))
    box(draw, (300, 760, 790, 1000), "诊断重点", "先确认下电、验电与绝缘安全\n再检查快充口锁止、HVIL 回路、BMS 充电允许和充电 CAN 通讯", h_font, b_font, accent="#DC2626")
    box(draw, (930, 760, 1400, 1000), "证据约束", "引用准确率 1.0000\n避免直接建议更换快充座\n资料不足时提示补充检测数据", h_font, b_font, accent="#16A34A")
    box(draw, (1540, 760, 2030, 1000), "论文价值", "展示从问题解析到证据定位、生成和复核的闭环\n支撑端到端案例分析", h_font, b_font, accent="#7C3AED")
    img.save(path)
    return path


def normalize_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3)
    section.right_margin = Cm(2)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)

    styles = doc.styles
    configure_style(styles["Normal"], "宋体", 14, bold=False)
    for style_name in ["Heading 1", "标题 1"]:
        if style_name in styles:
            configure_style(styles[style_name], "宋体", 15, bold=True)
    for style_name in ["Heading 2", "标题 2"]:
        if style_name in styles:
            configure_style(styles[style_name], "宋体", 14, bold=True)

    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        if p.style.name.startswith("Heading") or is_chapter_heading(text):
            p.paragraph_format.first_line_indent = None
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(8)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if is_top_heading(text) else WD_ALIGN_PARAGRAPH.LEFT
        elif text.startswith("图") and " " in text[:6]:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = None
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)
            set_paragraph_font(p, size=10.5)
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.first_line_indent = Cm(0.85)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
            p.paragraph_format.space_after = Pt(0)
        set_paragraph_font(p)

    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        for ri, row in enumerate(table.rows):
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                set_cell_margins(cell, top=100, start=120, bottom=100, end=120)
                for p in cell.paragraphs:
                    p.paragraph_format.first_line_indent = None
                    p.paragraph_format.line_spacing = 1.15
                    set_paragraph_font(p, size=10.5, bold=ri == 0)
                    if ri == 0:
                        shade_cell(cell, "E2E8F0")


def configure_style(style, east_asia: str, size: float, bold: bool) -> None:
    font = style.font
    font.name = "Times New Roman"
    font.size = Pt(size)
    font.bold = bold
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), east_asia)
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")


def set_paragraph_font(paragraph, size: float | None = None, bold: bool | None = None) -> None:
    for run in paragraph.runs:
        set_run_font(run, size=size, bold=bold)


def set_run_font(run, size: float | None = None, bold: bool | None = None) -> None:
    run.font.name = "Times New Roman"
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), "宋体")
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")


def set_cell_margins(cell, top=80, start=80, bottom=80, end=80):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)


def is_chapter_heading(text: str) -> bool:
    return bool(text and (text[0].isdigit() or text in {"摘要", "Abstract", "目录", "参考文献", "致谢"} or text.startswith("附录")))


def is_top_heading(text: str) -> bool:
    if text in {"摘要", "Abstract", "目录", "参考文献", "致谢"} or text.startswith("附录"):
        return True
    return bool(len(text) >= 3 and text[0].isdigit() and " " in text[:3] and "." not in text.split()[0])


def clear_paragraph(paragraph) -> None:
    for child in list(paragraph._p):
        paragraph._p.remove(child)


def set_para_text(paragraph, text: str, *, style: str | None = None, size: float | None = None, bold: bool | None = None) -> None:
    clear_paragraph(paragraph)
    if style:
        paragraph.style = style
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold)


def remove_element(element) -> None:
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def remove_internal_tables(doc: Document) -> None:
    for table in list(doc.tables):
        text = table_text(table)
        if text.startswith("项目内容论文定位") or text.startswith("使用说明本文档是基于当前代码仓库"):
            remove_element(table._element)


def replace_paragraphs(doc: Document) -> None:
    for p in list(doc.paragraphs):
        text = p.text.strip()
        if "多模态检索增强问答系统设计与实现" in text and "新能源汽车故障诊断" in text:
            set_para_text(p, TITLE, size=15, bold=True)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif text in PARAGRAPH_REPLACEMENTS:
            set_para_text(p, PARAGRAPH_REPLACEMENTS[text])
        elif text in REPLACEMENTS:
            set_para_text(p, REPLACEMENTS[text])
        elif text.startswith("新能源汽车售后维修场景具有知识来源分散"):
            set_para_text(p, CN_ABSTRACT)
        elif text.startswith("系统采用 Next.js 前端"):
            remove_element(p._element)
        elif text.startswith("实验部分基于 64 条覆盖故障码"):
            remove_element(p._element)
        elif text.startswith("Maintenance diagnosis for new energy vehicles is challenged"):
            set_para_text(p, EN_ABSTRACT)
        elif text.startswith("The system is implemented with Next.js"):
            remove_element(p._element)
        elif text.startswith("Experiments are conducted on 64 evaluation"):
            remove_element(p._element)
        elif text == "附录 B 实验复现实施说明":
            set_para_text(p, text)
            insert_toc_acknowledgement_after(doc, p)
        else:
            for old, new in REPLACEMENTS.items():
                if old in text:
                    set_para_text(p, text.replace(old, new))
                    break


def rebuild_toc(doc: Document) -> None:
    paragraphs = doc.paragraphs
    toc_idx = next((i for i, p in enumerate(paragraphs) if p.text.strip() == "目录"), None)
    body_idx = next(
        (
            i
            for i, p in enumerate(paragraphs)
            if p.text.strip() == "1 绪论" and ("Heading" in p.style.name or "标题" in p.style.name)
        ),
        None,
    )
    if toc_idx is None or body_idx is None or body_idx <= toc_idx:
        return
    body = doc.element.body
    for p in paragraphs[toc_idx + 1 : body_idx]:
        remove_element(p._element)
    insert_index = list(body).index(doc.paragraphs[toc_idx]._element) + 1
    for line in reversed(TOC_LINES):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.first_line_indent = None
        p.paragraph_format.left_indent = Cm(3.8)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(line)
        set_run_font(run, size=14)
        body.remove(p._element)
        body.insert(insert_index, p._element)


def replace_placeholder_figures(doc: Document, figure_paths: dict[str, Path]) -> None:
    for table in list(doc.tables):
        text = table_text(table)
        if "【需补图】" not in text:
            continue
        fig_key = next((key for key in figure_paths if key in text), None)
        if not fig_key:
            continue
        insert_picture_before_table(doc, table, figure_paths[fig_key], width=Inches(6.2))
        remove_element(table._element)


def insert_picture_before_table(doc: Document, table, image_path: Path, width):
    body = doc.element.body
    index = list(body).index(table._element)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = None
    p.add_run().add_picture(str(image_path), width=width)
    body.remove(p._element)
    body.insert(index, p._element)


def rewrite_supporting_tables(doc: Document) -> None:
    for table in doc.tables:
        text = table_text(table)
        if "待补测" in text and "答案正确率" in text:
            rewrite_table(
                table,
                [
                    ["评分维度", "记录方式", "判定标准"],
                    ["答案正确", "0/1", "是否覆盖主要根因和关键排查步骤"],
                    ["引用匹配", "0/1", "引用片段是否真实支持答案"],
                    ["安全合规", "0/1", "高压、涉水、绝缘等场景是否给出必要安全提醒"],
                    ["无幻觉", "0/1", "是否避免编造资料中没有的标准值、部件或结论"],
                    ["步骤可执行", "0/1", "是否给出可实际操作的检测动作、工具或标准值"],
                ],
            )
        elif text.startswith("编号位置需要补的图"):
            rewrite_table(
                table,
                [
                    ["编号", "位置", "图表", "来源或复现方式"],
                    ["A1", "3.3", "系统总体架构图", "依据 README、schema.py、rag/service.py 绘制"],
                    ["A2", "3.5", "模型端点与任务路由拓扑图", "依据 model_endpoints、model_routes 和 docs/model-config.md 绘制"],
                    ["A3", "3.5", "学校服务器 Docker Compose 部署拓扑图", "依据 docker-compose.gpu.yml 与 deploy/ 脚本绘制"],
                    ["A4", "4.2", "多模态 RAG 问答流程图", "依据 rag/service.py 与 provider.py 绘制"],
                    ["A5", "4.6", "系统主要功能界面示意图", "依据 apps/web 页面结构绘制，可在最终版补真实截图"],
                    ["A6", "5.1-5.4", "检索和知识库实验图", "来自 docs/paper/figures 与 apps/api/experiments/paper_experiments.py"],
                    ["A7", "5.5", "快充互锁异常案例流程图", "依据检索结果和案例分析文字绘制"],
                ],
            )


def rewrite_table(table, rows: list[list[str]]) -> None:
    while len(table.rows) < len(rows):
        table.add_row()
    while len(table.rows) > len(rows):
        remove_element(table.rows[-1]._tr)
    col_count = max(len(r) for r in rows)
    while len(table.columns) < col_count:
        add_column(table)
    for r, values in enumerate(rows):
        for c in range(col_count):
            cell = table.cell(r, c)
            cell.text = values[c] if c < len(values) else ""
            for p in cell.paragraphs:
                set_paragraph_font(p, size=10.5, bold=r == 0)
                p.paragraph_format.first_line_indent = None
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if r == 0 else WD_ALIGN_PARAGRAPH.LEFT
            if r == 0:
                shade_cell(cell, "E2E8F0")


def add_column(table):
    for row in table.rows:
        tc = OxmlElement("w:tc")
        tcPr = OxmlElement("w:tcPr")
        tc.append(tcPr)
        p = OxmlElement("w:p")
        tc.append(p)
        row._tr.append(tc)


def add_table_captions(doc: Document) -> None:
    seen: set[str] = set()
    command_count = 0
    for table in list(doc.tables):
        first = table.cell(0, 0).text.strip().replace("\n", "")
        caption = TABLE_CAPTIONS.get(first)
        if first == "命令":
            command_count += 1
            caption = f"表B-{command_count} " + ("检索评测复现命令" if command_count == 1 else "论文图表生成命令")
        if not caption:
            continue
        if caption in seen:
            continue
        seen.add(caption)
        insert_caption_before_table(doc, table, caption)


def insert_caption_before_table(doc: Document, table, caption: str):
    body = doc.element.body
    index = list(body).index(table._element)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = None
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(caption)
    set_run_font(run, size=10.5)
    body.remove(p._element)
    body.insert(index, p._element)


def rewrite_references(doc: Document) -> None:
    candidates = []
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() == "参考文献":
            candidates.append(i)
    start = candidates[-1] if candidates else None
    end = None
    if start is None:
        return
    for i, p in enumerate(doc.paragraphs[start + 1 :], start + 1):
        text = p.text.strip()
        if text.startswith("附录 A") or text.startswith("附录A"):
            end = i
            break
    if end is None:
        return
    for p in doc.paragraphs[start + 1:end]:
        remove_element(p._element)
    anchor = doc.paragraphs[start]
    body = doc.element.body
    insert_index = list(body).index(anchor._element) + 1
    for ref in reversed(REFERENCES):
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = None
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(ref)
        set_run_font(run, size=12)
        body.remove(p._element)
        body.insert(insert_index, p._element)


def insert_toc_acknowledgement_after(doc: Document, paragraph) -> None:
    body = doc.element.body
    items = [p.text.strip() for p in doc.paragraphs]
    if "致谢" in items[: max(0, items.index("1 绪论") if "1 绪论" in items else 0)]:
        return
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = None
    p.add_run("致谢")
    set_paragraph_font(p)
    insert_index = list(body).index(paragraph._element) + 1
    body.remove(p._element)
    body.insert(insert_index, p._element)


def add_acknowledgement(doc: Document) -> None:
    body = doc.element.body
    chapter_start = next((i for i, p in enumerate(doc.paragraphs) if p.text.strip() == "1 绪论"), 10)
    for idx, p in list(enumerate(doc.paragraphs)):
        text = p.text.strip()
        if idx > chapter_start and (text == "致谢" or "感谢指导教师在选题定位" in text):
            remove_element(p._element)

    heading = doc.add_paragraph()
    heading.style = doc.styles["Heading 1"]
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.paragraph_format.first_line_indent = None
    run = heading.add_run("致谢")
    set_run_font(run, size=15, bold=True)

    paragraph = doc.add_paragraph(
        "感谢指导教师在选题定位、系统设计、实验组织和论文撰写过程中给予的指导。感谢学院提供毕业设计研究与答辩环境，感谢同学在系统测试、演示反馈和材料整理方面提供的帮助。"
        "本文仍有不足，后续将继续围绕真实维修数据接入、工单闭环和模型推理效率优化开展改进。"
    )
    paragraph.paragraph_format.first_line_indent = Cm(0.85)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    set_paragraph_font(paragraph)

    sect = body[-1]
    body.remove(heading._element)
    body.remove(paragraph._element)
    body.insert(len(body) - 1, heading._element)
    body.insert(len(body) - 1, paragraph._element)


def configure_footer_page_numbers(doc: Document) -> None:
    for section in doc.sections:
        footer = section.footer
        while len(footer.paragraphs) > 1:
            remove_element(footer.paragraphs[-1]._element)
        p = footer.paragraphs[0]
        clear_paragraph(p)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = None
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), "PAGE")
        r = OxmlElement("w:r")
        rPr = OxmlElement("w:rPr")
        rFonts = OxmlElement("w:rFonts")
        rFonts.set(qn("w:ascii"), "Times New Roman")
        rFonts.set(qn("w:hAnsi"), "Times New Roman")
        rFonts.set(qn("w:eastAsia"), "宋体")
        rPr.append(rFonts)
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), "21")
        rPr.append(sz)
        r.append(rPr)
        t = OxmlElement("w:t")
        t.text = "1"
        r.append(t)
        fld.append(r)
        p._p.append(fld)


def enforce_page_breaks(doc: Document) -> None:
    first_body = next(
        (
            p
            for p in doc.paragraphs
            if p.text.strip() == "1 绪论" and ("Heading" in p.style.name or "标题" in p.style.name)
        ),
        None,
    )
    if first_body is not None:
        first_body.paragraph_format.page_break_before = True


def table_text(table) -> str:
    return "".join(cell.text.replace("\n", "") for row in table.rows for cell in row.cells).strip()


if __name__ == "__main__":
    main()
