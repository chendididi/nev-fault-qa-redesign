from io import BytesIO, StringIO
import csv
from pypdf import PdfReader
from docx import Document as DocxDocument
from openpyxl import load_workbook
from app.models.provider import OpenAICompatibleProvider


async def extract_text(filename: str, content_type: str, data: bytes, ocr_provider: OpenAICompatibleProvider | None = None) -> str:
    lower = filename.lower()
    if lower.endswith(".pdf"):
        reader = PdfReader(BytesIO(data))
        page_texts = [page.extract_text() or "" for page in reader.pages]
        if ocr_provider and any(not page_text.strip() for page_text in page_texts):
            page_texts = await _ocr_empty_pdf_pages(filename, data, page_texts, ocr_provider)
        return "\n".join(page_texts)
    if lower.endswith(".docx"):
        doc = DocxDocument(BytesIO(data))
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    if lower.endswith((".xlsx", ".xlsm")):
        wb = load_workbook(BytesIO(data), data_only=True, read_only=True)
        rows = []
        for sheet in wb.worksheets:
            rows.append(f"工作表: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                values = [str(value) for value in row if value is not None]
                if values:
                    rows.append(" | ".join(values))
        return "\n".join(rows)
    if lower.endswith(".csv"):
        text = data.decode("utf-8", errors="ignore")
        rows = csv.reader(StringIO(text))
        return "\n".join(" | ".join(row) for row in rows)
    if lower.endswith((".png", ".jpg", ".jpeg", ".webp", ".bmp")) or content_type.startswith("image/"):
        if not ocr_provider:
            raise RuntimeError("图片 OCR 需要可用的 gpt-5.5 视觉接口")
        return (await ocr_provider.ocr_image(data, content_type or "image/png", context=filename)).strip()
    return data.decode("utf-8", errors="ignore")


async def _ocr_empty_pdf_pages(
    filename: str,
    data: bytes,
    page_texts: list[str],
    ocr_provider: OpenAICompatibleProvider,
) -> list[str]:
    try:
        import fitz
    except ImportError as exc:
        raise RuntimeError("PDF 图片页 OCR 需要安装 PyMuPDF") from exc

    doc = fitz.open(stream=data, filetype="pdf")
    try:
        result = list(page_texts)
        for index, page_text in enumerate(page_texts):
            if page_text.strip():
                continue
            page = doc.load_page(index)
            pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
            image_bytes = pixmap.tobytes("png")
            result[index] = await ocr_provider.ocr_image(
                image_bytes,
                "image/png",
                context=f"{filename} 第 {index + 1} 页",
            )
        return result
    finally:
        doc.close()


def chunk_text(text: str, *, size: int = 900, overlap: int = 120) -> list[str]:
    normalized = "\n".join(line.strip() for line in text.splitlines() if line.strip())
    if not normalized:
        return []
    chunks: list[str] = []
    start = 0
    while start < len(normalized):
        end = min(start + size, len(normalized))
        chunks.append(normalized[start:end])
        if end == len(normalized):
            break
        start = max(0, end - overlap)
    return chunks
